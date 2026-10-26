"""Apply deterministic Word layout after the professor's Quarto render step."""
from pathlib import Path
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

path=Path(sys.argv[1] if len(sys.argv)>1 else 'replication-lab.docx')
doc=Document(path)
for section in doc.sections:
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
for paragraph in doc.paragraphs:
    if paragraph.style.name in ['Title','Subtitle'] or paragraph.style.name.startswith('Heading'):
        for run in paragraph.runs:run.font.color.rgb=RGBColor(0,0,0)
    if paragraph._p.xpath('.//w:drawing'):paragraph.paragraph_format.keep_with_next=True
    if paragraph.style.name == 'Bibliography':paragraph.paragraph_format.keep_together=True
for shape in doc.inline_shapes:
    if shape.width > Inches(6.25):
        ratio=Inches(6.25)/shape.width
        shape.height=int(shape.height*ratio);shape.width=Inches(6.25)
for name in ['Source Code','Verbatim Char']:
    if name in doc.styles:
        doc.styles[name].font.name='Consolas';doc.styles[name].font.size=Pt(9)
for table in doc.tables:
    table.autofit=False
    n=len(table.columns)
    fractions=([.25]+[(1-.25)/(n-1)]*(n-1)) if n>1 else [1]
    for i,column in enumerate(table.columns):column.width=Inches(6.5*fractions[i])
    borders=table._tbl.tblPr.find(qn('w:tblBorders'))
    if borders is None:
        borders=OxmlElement('w:tblBorders');table._tbl.tblPr.append(borders)
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    for ri,row in enumerate(table.rows):
        props=row._tr.get_or_add_trPr()
        if ri==0:
            repeat=OxmlElement('w:tblHeader');props.append(repeat)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(6.5*fractions[ci]);cell.vertical_alignment=1
            tcpr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            if ri==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E8EFF3');tcpr.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before=Pt(1);paragraph.paragraph_format.space_after=Pt(1)
                paragraph.paragraph_format.line_spacing=1.0;paragraph.alignment=0 if ci==0 else 1
                for run in paragraph.runs:
                    run.font.name='Calibri';run.font.size=Pt(9.5);run.font.color.rgb=RGBColor(0,0,0)
                    if ri==0:run.bold=True
doc.core_properties.author='Hemendra Gummadi';doc.core_properties.title='When the Test Environment Changes'
doc.save(path)
print('Formatted',path)
